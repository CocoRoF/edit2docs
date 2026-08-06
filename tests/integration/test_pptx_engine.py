"""Structural PPTX editor (documents/pptx_engine): addressable outline +
surgical, byte-preserving operations on real decks (text / table / chart).

These are the guarantees the legacy SVG-rewrite path could not give: tables
and charts are edited IN PLACE (not flattened), and everything an edit does
not touch stays byte-identical.
"""

from __future__ import annotations

import io
import zipfile

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Emu, Inches, Pt

from edit2docs.documents.pptx_engine import (
    OP_CATALOG,
    PptxEdit,
    apply_pptx_edits,
    describe_ops,
    find_shapes,
    pptx_outline,
)


def _deck() -> bytes:
    """Slide 1: title textbox + 3x3 table + column chart. Slide 2: text only."""
    prs = Presentation()
    prs.slide_width = Emu(12192000)
    prs.slide_height = Emu(6858000)

    s1 = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s1.shapes.add_textbox(Inches(1), Inches(0.3), Inches(6), Inches(1))
    tb.text_frame.text = "원본 제목"
    tbl = s1.shapes.add_table(3, 3, Inches(0.5), Inches(1.5), Inches(6), Inches(2)).table
    for r in range(3):
        for c in range(3):
            tbl.cell(r, c).text = f"R{r}C{c}"
    cd = CategoryChartData()
    cd.categories = ["Q1", "Q2", "Q3"]
    cd.add_series("매출", (10.0, 20.0, 30.0))
    s1.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(3.8), Inches(5), Inches(3), cd
    )

    s2 = prs.slides.add_slide(prs.slide_layouts[6])
    b2 = s2.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    b2.text_frame.text = "둘째 슬라이드"

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _ids(deck: bytes):
    outline = pptx_outline(deck)
    title = next(e for e in outline if e.get("kind") == "text" and e["slide"] == 1)
    table = next(e for e in outline if e.get("kind") == "table")
    chart = next(e for e in outline if e.get("kind") == "chart")
    return title["shape"], table["shape"], chart["chart"]


def _members(pptx: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(pptx)) as z:
        return {n: z.read(n) for n in z.namelist()}


class TestOutline:
    def test_outline_lists_text_table_and_chart(self):
        outline = pptx_outline(_deck())
        kinds = {e.get("kind") for e in outline}
        assert {"text", "table", "chart"} <= kinds
        # table cells are addressable
        cells = [e for e in outline if "table" in e and "row" in e]
        assert len(cells) == 9
        assert {(e["row"], e["col"]) for e in cells} == {(r, c) for r in range(3) for c in range(3)}
        # chart series data is surfaced
        chart = next(e for e in outline if e.get("kind") == "chart")
        assert chart["series"][0]["categories"] == ["Q1", "Q2", "Q3"]
        assert chart["series"][0]["values"] == [10.0, 20.0, 30.0]


class TestSurgicalEdits:
    def test_all_op_types_apply(self):
        deck = _deck()
        title, table, _ = _ids(deck)
        edits = [
            PptxEdit("set_text", slide=1, shape=title, para=0, new_text="새 제목", old_text="원본 제목"),
            PptxEdit("set_table_cell", slide=1, shape=table, row=0, col=0, new_text="바뀐셀", old_text="R0C0"),
            PptxEdit("insert_row", slide=1, shape=table, at=3),
            PptxEdit("set_table_cell", slide=1, shape=table, row=3, col=0, new_text="새행"),
            PptxEdit("set_chart_data", slide=1, chart=0,
                     categories=["Q1", "Q2", "Q3", "Q4"], series=[{"name": "매출", "values": [11, 22, 33, 44]}]),
            PptxEdit("set_chart_title", slide=1, chart=0, title="분기 매출"),
        ]
        out, results = apply_pptx_edits(deck, edits)
        assert [r.status for r in results] == ["applied"] * 6

        prs = Presentation(io.BytesIO(out))
        s1 = prs.slides[0]
        texts = [sh.text_frame.text for sh in s1.shapes if sh.has_text_frame]
        assert "새 제목" in texts
        tbl = next(sh.table for sh in s1.shapes if sh.has_table)
        assert len(tbl.rows) == 4
        assert tbl.cell(0, 0).text == "바뀐셀"
        assert tbl.cell(3, 0).text == "새행"
        chart = next(sh.chart for sh in s1.shapes if sh.has_chart)
        assert chart.has_title and chart.chart_title.text_frame.text == "분기 매출"
        assert list(chart.plots[0].categories) == ["Q1", "Q2", "Q3", "Q4"]
        assert list(chart.series[0].values) == [11.0, 22.0, 33.0, 44.0]

    def test_untouched_parts_are_byte_identical(self):
        """Editing slide 1 must not rewrite slide 2 (or any unrelated part)."""
        deck = _deck()
        title, _, _ = _ids(deck)
        before = _members(deck)
        out, results = apply_pptx_edits(
            deck, [PptxEdit("set_text", slide=1, shape=title, para=0, new_text="X")]
        )
        assert results[0].status == "applied"
        after = _members(out)
        # slide2.xml is untouched → byte-identical
        assert after["ppt/slides/slide2.xml"] == before["ppt/slides/slide2.xml"]
        # the theme / presentation parts are untouched too
        assert after["ppt/theme/theme1.xml"] == before["ppt/theme/theme1.xml"]
        # slide1 DID change
        assert after["ppt/slides/slide1.xml"] != before["ppt/slides/slide1.xml"]

    def test_editing_a_chart_preserves_other_slides(self):
        """set_chart_data rewrites the chart part + embedded workbook — but not
        an unrelated slide."""
        deck = _deck()
        before = _members(deck)
        out, results = apply_pptx_edits(deck, [
            PptxEdit("set_chart_data", slide=1, chart=0,
                     categories=["Q1", "Q2"], series=[{"name": "s", "values": [5, 6]}]),
        ])
        assert results[0].status == "applied"
        after = _members(out)
        assert after["ppt/slides/slide2.xml"] == before["ppt/slides/slide2.xml"]
        chart = next(sh.chart for sh in Presentation(io.BytesIO(out)).slides[0].shapes if sh.has_chart)
        assert list(chart.series[0].values) == [5.0, 6.0]

    def test_addresses_target_the_right_slide(self):
        """An op on slide 2 must not touch slide 1 (multi-slide addressing)."""
        deck = _deck()
        s2_text = next(
            e for e in pptx_outline(deck) if e.get("kind") == "text" and e["slide"] == 2
        )
        out, results = apply_pptx_edits(deck, [
            PptxEdit("set_text", slide=2, shape=s2_text["shape"], para=0, new_text="둘째 변경"),
        ])
        assert results[0].status == "applied"
        prs = Presentation(io.BytesIO(out))
        assert "둘째 변경" in [sh.text_frame.text for sh in prs.slides[1].shapes if sh.has_text_frame]
        # slide 1's title untouched
        assert "원본 제목" in [sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame]

    def test_delete_row(self):
        deck = _deck()
        _, table, _ = _ids(deck)
        out, results = apply_pptx_edits(deck, [PptxEdit("delete_row", slide=1, shape=table, row=1)])
        assert results[0].status == "applied"
        tbl = next(sh.table for sh in Presentation(io.BytesIO(out)).slides[0].shapes if sh.has_table)
        assert len(tbl.rows) == 2
        assert tbl.cell(1, 0).text == "R2C0"  # row 1 gone, old row 2 shifted up


class TestGoldenDeckPreservation:
    """The core guarantee the SVG rewrite could not give: editing some content
    NEVER drops or alters the styled header, theme, or native objects the edit
    did not touch. This is the regression for '스타일 헤더가 맘대로 지워짐'."""

    def _styled_deck(self) -> tuple[bytes, int, int, int]:
        prs = Presentation()
        prs.slide_width = Emu(12192000)
        prs.slide_height = Emu(6858000)
        s = prs.slides.add_slide(prs.slide_layouts[6])
        # A styled "header" band we will NOT touch.
        hdr = s.shapes.add_textbox(Inches(0), Inches(0), Inches(13.3), Inches(0.8))
        hdr.text_frame.text = "COMPANY HEADER"
        run = hdr.text_frame.paragraphs[0].runs[0]
        run.font.bold = True
        run.font.size = Pt(20)
        from pptx.dml.color import RGBColor
        run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
        # A title we WILL edit.
        title = s.shapes.add_textbox(Inches(1), Inches(1.2), Inches(6), Inches(1))
        title.text_frame.text = "원본 제목"
        # A table we WILL edit one cell of.
        tbl = s.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(5), Inches(1.5)).table
        for r in range(2):
            for c in range(2):
                tbl.cell(r, c).text = f"R{r}C{c}"
        buf = io.BytesIO()
        prs.save(buf)
        deck = buf.getvalue()
        return deck, hdr.shape_id, title.shape_id, tbl_shape_id(deck)

    def test_untouched_styled_header_survives_edits_to_other_shapes(self):
        deck, hdr_id, title_id, table_id = self._styled_deck()
        out, results = apply_pptx_edits(deck, [
            PptxEdit("set_text", slide=1, shape=title_id, para=0, new_text="새 제목"),
            PptxEdit("set_table_cell", slide=1, shape=table_id, row=0, col=0, new_text="변경"),
        ])
        assert [r.status for r in results] == ["applied", "applied"]

        prs = Presentation(io.BytesIO(out))
        shapes = {sh.shape_id: sh for sh in prs.slides[0].shapes}
        # The header is byte-perfectly intact: text, bold, size, color.
        hdr = shapes[hdr_id]
        hrun = hdr.text_frame.paragraphs[0].runs[0]
        assert hdr.text_frame.text == "COMPANY HEADER"
        assert hrun.font.bold is True
        assert hrun.font.size == Pt(20)
        assert str(hrun.font.color.rgb) == "C00000"
        # The targeted edits landed.
        assert shapes[title_id].text_frame.text == "새 제목"
        tbl = next(sh.table for sh in prs.slides[0].shapes if sh.has_table)
        assert tbl.cell(0, 0).text == "변경" and tbl.cell(1, 1).text == "R1C1"


def tbl_shape_id(deck: bytes) -> int:
    return next(e["shape"] for e in pptx_outline(deck) if e.get("kind") == "table")


class TestStyleAndPosition:
    def test_outline_reports_shape_geometry_in_inches(self):
        outline = pptx_outline(_deck())
        title = next(e for e in outline if e.get("kind") == "text" and e["slide"] == 1)
        assert title["pos"] == {"left": 1.0, "top": 0.3, "width": 6.0, "height": 1.0}

    def test_set_shape_style(self):
        deck = _deck()
        title, _, _ = _ids(deck)
        out, results = apply_pptx_edits(deck, [
            PptxEdit("set_shape_style", slide=1, shape=title,
                     color="FF0000", size_pt=28, bold=True, fill="EEEEEE"),
        ])
        assert results[0].status == "applied"
        sh = next(s for s in Presentation(io.BytesIO(out)).slides[0].shapes if s.has_text_frame)
        run = sh.text_frame.paragraphs[0].runs[0]
        assert str(run.font.color.rgb) == "FF0000"
        assert run.font.size == 355600  # 28pt
        assert run.font.bold is True
        assert str(sh.fill.fore_color.rgb) == "EEEEEE"
        # text is unchanged
        assert sh.text_frame.text == "원본 제목"

    def test_set_shape_position(self):
        deck = _deck()
        title, _, _ = _ids(deck)
        out, results = apply_pptx_edits(deck, [
            PptxEdit("set_shape_position", slide=1, shape=title,
                     left=2.0, top=3.0, width=6.0, height=1.5),
        ])
        assert results[0].status == "applied"
        sh = next(s for s in Presentation(io.BytesIO(out)).slides[0].shapes if s.has_text_frame)
        assert (round(sh.left / 914400, 2), round(sh.top / 914400, 2),
                round(sh.width / 914400, 2), round(sh.height / 914400, 2)) == (2.0, 3.0, 6.0, 1.5)

    def test_set_cell_style(self):
        deck = _deck()
        _, table, _ = _ids(deck)
        out, results = apply_pptx_edits(deck, [
            PptxEdit("set_cell_style", slide=1, shape=table, row=0, col=0,
                     fill="0000FF", color="FFFFFF", bold=True),
        ])
        assert results[0].status == "applied"
        cell = next(
            sh.table for sh in Presentation(io.BytesIO(out)).slides[0].shapes if sh.has_table
        ).cell(0, 0)
        assert cell.text == "R0C0"  # text unchanged
        assert str(cell.fill.fore_color.rgb) == "0000FF"
        assert str(cell.text_frame.paragraphs[0].runs[0].font.color.rgb) == "FFFFFF"

    def test_style_needs_an_attribute(self):
        deck = _deck()
        title, _, _ = _ids(deck)
        _, results = apply_pptx_edits(deck, [PptxEdit("set_shape_style", slide=1, shape=title)])
        assert results[0].status == "invalid"


class TestReviewFixes:
    """Regressions for the adversarial-review findings."""

    def test_set_shape_style_para0_targets_only_first_paragraph(self):
        prs = Presentation()
        prs.slide_width = Emu(12192000)
        prs.slide_height = Emu(6858000)
        s = prs.slides.add_slide(prs.slide_layouts[6])
        tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
        tb.text_frame.paragraphs[0].add_run().text = "First"
        tb.text_frame.add_paragraph().add_run().text = "Second"
        sid = tb.shape_id
        buf = io.BytesIO()
        prs.save(buf)
        deck = buf.getvalue()

        out, _ = apply_pptx_edits(deck, [PptxEdit("set_shape_style", slide=1, shape=sid, para=0, bold=True)])
        tf = next(x for x in Presentation(io.BytesIO(out)).slides[0].shapes if x.has_text_frame).text_frame
        assert tf.paragraphs[0].runs[0].font.bold is True
        assert tf.paragraphs[1].runs[0].font.bold is None  # para 1 untouched

        out2, _ = apply_pptx_edits(deck, [PptxEdit("set_shape_style", slide=1, shape=sid, bold=True)])
        tf2 = next(x for x in Presentation(io.BytesIO(out2)).slides[0].shapes if x.has_text_frame).text_frame
        assert tf2.paragraphs[0].runs[0].font.bold is True
        assert tf2.paragraphs[1].runs[0].font.bold is True  # no para = whole shape

    def test_insert_row_then_populate_is_order_independent(self):
        deck = _deck()
        _, table, _ = _ids(deck)
        for order in ("insert_first", "cell_first"):
            ops = (
                [PptxEdit("insert_row", slide=1, shape=table, at=1),
                 PptxEdit("set_table_cell", slide=1, shape=table, row=1, col=0, new_text="NEW")]
                if order == "insert_first" else
                [PptxEdit("set_table_cell", slide=1, shape=table, row=1, col=0, new_text="NEW"),
                 PptxEdit("insert_row", slide=1, shape=table, at=1)]
            )
            out, _ = apply_pptx_edits(deck, ops)
            tbl = next(sh.table for sh in Presentation(io.BytesIO(out)).slides[0].shapes if sh.has_table)
            assert len(tbl.rows) == 4
            assert tbl.cell(1, 0).text == "NEW"       # new row populated
            assert tbl.cell(2, 0).text == "R1C0"      # old row 1 shifted down

    def test_literal_newline_in_run_keeps_para_index_aligned(self):
        prs = Presentation()
        prs.slide_width = Emu(12192000)
        prs.slide_height = Emu(6858000)
        s = prs.slides.add_slide(prs.slide_layouts[6])
        tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
        tb.text_frame.paragraphs[0].add_run().text = "Hello\nWorld"  # newline in one run
        tb.text_frame.add_paragraph().add_run().text = "Bye"
        sid = tb.shape_id
        buf = io.BytesIO()
        prs.save(buf)
        deck = buf.getvalue()
        paras = [e for e in pptx_outline(deck) if e.get("kind") == "text"]
        assert len(paras) == 2  # two a:p, not three (would be 3 if split on "\n")
        out, results = apply_pptx_edits(deck, [PptxEdit("set_text", slide=1, shape=sid, para=1, new_text="EDITED")])
        assert results[0].status == "applied"
        tf = next(x for x in Presentation(io.BytesIO(out)).slides[0].shapes if x.has_text_frame).text_frame
        assert tf.paragraphs[0].text == "Hello\nWorld"  # untouched
        assert tf.paragraphs[1].text == "EDITED"        # the right a:p

    def test_string_addresses_never_crash_the_batch(self):
        """Even a direct caller passing string addresses degrades to a per-op
        result instead of raising out of apply_pptx_edits."""
        deck = _deck()
        _, table, _ = _ids(deck)
        _, results = apply_pptx_edits(deck, [
            PptxEdit("set_text", slide="1", shape="9999", new_text="x"),   # bad shape
            PptxEdit("set_table_cell", slide="1", shape=str(table), row="0", col="0", new_text="Y"),
        ])
        assert results[0].status == "not_found"
        assert results[1].status == "applied"  # string table id + row/col coerced


class TestFrameworkInterface:
    def test_catalog_is_self_describing_and_matches_actions(self):
        cat = describe_ops()
        assert cat is not OP_CATALOG  # a copy, not the live dict
        # every catalog action is valid, and each has address+payload contracts
        for action, spec in cat.items():
            assert "address" in spec and "payload" in spec and "summary" in spec

    def test_find_shapes_by_text_and_kind(self):
        deck = _deck()
        by_text = find_shapes(deck, text="원본")
        assert any(e["text"] == "원본 제목" for e in by_text)
        assert all(e.get("kind") == "chart" for e in find_shapes(deck, kind="chart"))
        assert find_shapes(deck, slide=2) and all(e["slide"] == 2 for e in find_shapes(deck, slide=2))

    def test_dry_run_validates_without_writing(self):
        deck = _deck()
        title, _, _ = _ids(deck)
        out, results = apply_pptx_edits(
            deck, [PptxEdit("set_text", slide=1, shape=title, new_text="X")], dry_run=True
        )
        assert results[0].status == "applied"  # would apply
        assert out == deck                       # but nothing written

    def test_atomic_rolls_back_the_whole_batch_on_any_failure(self):
        deck = _deck()
        title, _, _ = _ids(deck)
        out, results = apply_pptx_edits(deck, [
            PptxEdit("set_text", slide=1, shape=title, new_text="APPLIED"),
            PptxEdit("set_text", slide=9, shape=1, new_text="NOPE"),  # not_found
        ], atomic=True)
        assert [r.status for r in results] == ["applied", "not_found"]
        assert out == deck  # the applied op was rolled back too


class TestAuthoringOps:
    def test_set_runs_mixed_formatting(self):
        deck = _deck()
        title, _, _ = _ids(deck)
        out, results = apply_pptx_edits(deck, [
            PptxEdit("set_runs", slide=1, shape=title, para=0,
                     runs=[{"text": "Hello "}, {"text": "bold", "bold": True, "color": "FF0000"}]),
        ])
        assert results[0].status == "applied"
        runs = next(
            sh for sh in Presentation(io.BytesIO(out)).slides[0].shapes
            if sh.shape_id == title
        ).text_frame.paragraphs[0].runs
        assert [(r.text, r.font.bold) for r in runs] == [("Hello ", None), ("bold", True)]

    def test_add_delete_duplicate_shape(self):
        deck = _deck()
        title, _, _ = _ids(deck)
        n0 = len(Presentation(io.BytesIO(deck)).slides[0].shapes)
        out, results = apply_pptx_edits(deck, [
            PptxEdit("add_textbox", slide=1, new_text="new box",
                     left=1, top=6, width=3, height=1, bold=True),
            PptxEdit("duplicate_shape", slide=1, shape=title, left=7, top=1),
        ])
        assert [r.status for r in results] == ["applied", "applied"]
        s = Presentation(io.BytesIO(out)).slides[0]
        assert len(s.shapes) == n0 + 2
        assert "new box" in [sh.text_frame.text for sh in s.shapes if sh.has_text_frame]

        out2, r2 = apply_pptx_edits(deck, [PptxEdit("delete_shape", slide=1, shape=title)])
        assert r2[0].status == "applied"
        s2 = Presentation(io.BytesIO(out2)).slides[0]
        assert title not in {sh.shape_id for sh in s2.shapes}

    def test_table_insert_column_and_merge(self):
        deck = _deck()  # 3x3 table
        _, table, _ = _ids(deck)
        out, results = apply_pptx_edits(deck, [
            PptxEdit("insert_column", slide=1, shape=table, at=1),
            PptxEdit("set_table_cell", slide=1, shape=table, row=0, col=1, new_text="COL"),
            PptxEdit("merge_cells", slide=1, shape=table, row=0, col=0, row2=0, col2=1),
        ])
        assert [r.status for r in results] == ["applied"] * 3
        tbl = next(sh.table for sh in Presentation(io.BytesIO(out)).slides[0].shapes if sh.has_table)
        assert len(tbl.columns) == 4  # 3 + 1 inserted
        assert tbl.cell(0, 0)._tc.get("gridSpan") == "2"  # merged (0,0)-(0,1)

    def test_table_delete_column(self):
        deck = _deck()  # 3x3
        _, table, _ = _ids(deck)
        out, results = apply_pptx_edits(deck, [
            PptxEdit("delete_column", slide=1, shape=table, col=1),
        ])
        assert results[0].status == "applied"
        tbl = next(sh.table for sh in Presentation(io.BytesIO(out)).slides[0].shapes if sh.has_table)
        assert len(tbl.columns) == 2
        assert tbl.cell(0, 1).text == "R0C2"  # col 1 gone, old col 2 shifted left


class TestNativeAddSlide:
    def test_add_slide_from_layout_fills_placeholders(self):
        prs = Presentation()
        prs.slide_width = Emu(12192000)
        prs.slide_height = Emu(6858000)
        prs.slides.add_slide(prs.slide_layouts[6])  # one blank slide
        buf = io.BytesIO()
        prs.save(buf)
        deck = buf.getvalue()

        out, results = apply_pptx_edits(deck, [
            PptxEdit("add_slide", after=1, layout="Title and Content",
                     title="요약", body="한 줄\n두 줄\n세 줄"),
            PptxEdit("add_slide", layout="Section Header", title="2부"),
        ])
        assert [r.status for r in results] == ["applied", "applied"]
        prs2 = Presentation(io.BytesIO(out))
        assert len(prs2.slides) == 3
        s1 = prs2.slides[1]
        assert s1.slide_layout.name == "Title and Content"
        texts = {str(ph.placeholder_format.type): ph.text for ph in s1.placeholders if ph.text}
        assert any("요약" in v for v in texts.values())
        body = next(ph.text for ph in s1.placeholders if "\n" in (ph.text or ""))
        assert body.split("\n") == ["한 줄", "두 줄", "세 줄"]  # 3 real paragraphs
        assert prs2.slides[2].slide_layout.name == "Section Header"

    def test_add_slide_default_layout_is_title_and_body(self):
        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[6])
        buf = io.BytesIO()
        prs.save(buf)
        out, results = apply_pptx_edits(buf.getvalue(), [
            PptxEdit("add_slide", title="T", body="B"),
        ])
        assert results[0].status == "applied"
        s = Presentation(io.BytesIO(out)).slides[-1]
        types = {ph.placeholder_format.type for ph in s.placeholders}
        # a real content layout has both a title and a body-ish placeholder
        assert any("TITLE" in str(t) for t in types)


class TestCorrectnessR2:
    """Regressions for the second adversarial-review round."""

    def test_non_finite_and_ambiguous_chart_values_rejected(self):
        deck = _deck()
        _, results = apply_pptx_edits(deck, [
            PptxEdit("set_chart_data", slide=1, chart=0,
                     categories=["a", "b"], series=[{"name": "s", "values": [float("inf"), 1]}]),
            PptxEdit("set_chart_data", slide=1, chart=0,
                     categories=["a"], series=[{"name": "s", "values": ["1,5"]}]),
        ])
        assert results[0].status == "invalid"  # inf never written to numCache
        assert results[1].status == "invalid"  # ambiguous European decimal

    def test_thousands_separator_chart_value_ok(self):
        deck = _deck()
        out, results = apply_pptx_edits(deck, [
            PptxEdit("set_chart_data", slide=1, chart=0,
                     categories=["a"], series=[{"name": "s", "values": ["1,000"]}]),
        ])
        assert results[0].status == "applied"
        chart = next(sh.chart for sh in Presentation(io.BytesIO(out)).slides[0].shapes if sh.has_chart)
        assert list(chart.series[0].values) == [1000.0]

    def test_bool_address_is_rejected_not_coerced_to_one(self):
        deck = _deck()
        _, results = apply_pptx_edits(deck, [PptxEdit("set_text", slide=1, shape=True, new_text="x")])
        assert results[0].status in ("not_found", "invalid")  # not silently shape #1


class TestPhase2Ops:
    def _rich_deck(self):
        from pptx.chart.data import CategoryChartData
        from pptx.enum.chart import XL_CHART_TYPE
        prs = Presentation()
        prs.slide_width = Emu(12192000)
        prs.slide_height = Emu(6858000)
        s = prs.slides.add_slide(prs.slide_layouts[6])
        tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(2))
        tb.text_frame.paragraphs[0].add_run().text = "첫 줄"
        tb.text_frame.add_paragraph().add_run().text = "둘째 줄"
        cd = CategoryChartData()
        cd.categories = ["A", "B"]
        cd.add_series("S", (1.0, 2.0))
        s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(4), Inches(4), Inches(2), cd)
        buf = io.BytesIO()
        prs.save(buf)
        return buf.getvalue(), tb.shape_id

    def test_notes_bullets_hyperlink_zorder(self):
        deck, tid = self._rich_deck()
        out, results = apply_pptx_edits(deck, [
            PptxEdit("set_notes", slide=1, new_text="발표 노트\n둘째 줄"),
            PptxEdit("set_bullet", slide=1, shape=tid, para=0, bullet="bullet"),
            PptxEdit("set_bullet", slide=1, shape=tid, para=1, bullet="number"),
            PptxEdit("set_hyperlink", slide=1, shape=tid, para=0, run=0, url="https://x.com"),
            PptxEdit("set_z_order", slide=1, shape=tid, order="front"),
        ])
        assert [r.status for r in results] == ["applied"] * 5
        prs = Presentation(io.BytesIO(out))
        assert prs.slides[0].notes_slide.notes_text_frame.text.split("\n") == ["발표 노트", "둘째 줄"]
        run = next(sh for sh in prs.slides[0].shapes if sh.shape_id == tid).text_frame.paragraphs[0].runs[0]
        assert run.hyperlink.address == "https://x.com"

    def test_chart_legend_and_series_color(self):
        deck, _ = self._rich_deck()
        out, results = apply_pptx_edits(deck, [
            PptxEdit("set_legend", slide=1, chart=0, position="b"),
            PptxEdit("set_series_color", slide=1, chart=0, series_index=0, color="FF0000"),
        ])
        assert [r.status for r in results] == ["applied", "applied"]
        chart = next(sh.chart for sh in Presentation(io.BytesIO(out)).slides[0].shapes if sh.has_chart)
        assert chart.has_legend

    def test_theme_color_and_font(self):
        deck, _ = self._rich_deck()
        out, results = apply_pptx_edits(deck, [
            PptxEdit("set_theme_color", theme_name="accent1", color="00AA55"),
            PptxEdit("set_theme_font", which="minor", typeface="Malgun Gothic"),
        ])
        assert [r.status for r in results] == ["applied", "applied"]
        import zipfile
        th = zipfile.ZipFile(io.BytesIO(out)).read("ppt/theme/theme1.xml").decode()
        assert "00AA55" in th and "Malgun Gothic" in th

    def test_invalid_theme_slot_is_not_found(self):
        deck, _ = self._rich_deck()
        _, results = apply_pptx_edits(deck, [
            PptxEdit("set_theme_color", theme_name="nope", color="FFFFFF"),
        ])
        assert results[0].status == "not_found"


class TestGuardsAndErrors:
    def test_stale_old_text_is_rejected(self):
        deck = _deck()
        title, table, _ = _ids(deck)
        _, results = apply_pptx_edits(deck, [
            PptxEdit("set_text", slide=1, shape=title, para=0, new_text="X", old_text="틀린 원본"),
            PptxEdit("set_table_cell", slide=1, shape=table, row=0, col=0, new_text="X", old_text="wrong"),
        ])
        assert results[0].status == "stale"
        assert results[1].status == "stale"

    def test_not_found_addresses(self):
        deck = _deck()
        _, results = apply_pptx_edits(deck, [
            PptxEdit("set_text", slide=9, shape=2, new_text="X"),        # bad slide
            PptxEdit("set_text", slide=1, shape=9999, new_text="X"),     # bad shape
            PptxEdit("set_chart_data", slide=1, chart=5, categories=["a"], series=[{"name": "s", "values": [1]}]),
        ])
        assert results[0].status == "not_found"
        assert results[1].status == "not_found"
        assert results[2].status == "not_found"

    def test_ragged_chart_data_is_invalid(self):
        deck = _deck()
        _, results = apply_pptx_edits(deck, [
            PptxEdit("set_chart_data", slide=1, chart=0,
                     categories=["Q1", "Q2"], series=[{"name": "s", "values": [1, 2, 3]}]),
        ])
        assert results[0].status == "invalid"

    def test_unknown_action_is_invalid(self):
        deck = _deck()
        _, results = apply_pptx_edits(deck, [PptxEdit("frobnicate", slide=1, shape=2)])
        assert results[0].status == "invalid"

    def test_batch_is_order_preserving_in_results(self):
        deck = _deck()
        title, table, _ = _ids(deck)
        _, results = apply_pptx_edits(deck, [
            PptxEdit("delete_row", slide=1, shape=table, row=2),
            PptxEdit("set_text", slide=1, shape=title, para=0, new_text="Z"),
        ])
        # results correlate to input order even though row-op sorts internally
        assert results[0].action == "delete_row" and results[0].status == "applied"
        assert results[1].action == "set_text" and results[1].status == "applied"
