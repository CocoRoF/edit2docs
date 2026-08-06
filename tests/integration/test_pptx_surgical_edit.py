"""edit_document(fmt="pptx") — the surgical PPTX chat-edit path.

The planner LLM is stubbed; the addressable outline, op parsing, and the
byte-preserving apply run for real against a python-pptx-built deck with a
table and a chart. This is the path that replaces the lossy SVG rewrite for
targeted edits.
"""

from __future__ import annotations

import io

import pytest
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Emu, Inches

from edit2docs.llm.anthropic_client import LLMResult, LLMUsage
from edit2docs.tools.edit_doc import EditDocRequest, edit_document


def _deck() -> bytes:
    prs = Presentation()
    prs.slide_width = Emu(12192000)
    prs.slide_height = Emu(6858000)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Inches(1), Inches(0.3), Inches(6), Inches(1))
    tb.text_frame.text = "원본 제목"
    tbl = s.shapes.add_table(3, 3, Inches(0.5), Inches(1.5), Inches(6), Inches(2)).table
    for r in range(3):
        for c in range(3):
            tbl.cell(r, c).text = f"R{r}C{c}"
    cd = CategoryChartData()
    cd.categories = ["Q1", "Q2", "Q3"]
    cd.add_series("매출", (10.0, 20.0, 30.0))
    s.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(3.8), Inches(5), Inches(3), cd
    )
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


class _PlanLLM:
    """Returns a fixed plan; captures the outline the planner was shown."""

    def __init__(self, plan: str):
        self.plan = plan
        self.calls: list[str] = []

    async def complete(self, system_prompt, user_message, **kwargs):
        self.calls.append(user_message)
        return LLMResult(
            text=self.plan,
            usage=LLMUsage(input_tokens=10, output_tokens=10),
            model="stub",
            stop_reason="end_turn",
        )


PLAN = """```reply
1번 슬라이드의 제목·표 셀·차트 데이터를 수정합니다.
```
```edit_plan
operations:
  - action: set_text
    slide: 1
    shape: 2
    para: 0
    new_text: "새 제목"
    old_text: "원본 제목"
  - action: set_table_cell
    slide: 1
    shape: 3
    row: 0
    col: 0
    new_text: "바뀐셀"
  - action: set_chart_data
    slide: 1
    chart: 0
    categories: ["Q1", "Q2", "Q3", "Q4"]
    series:
      - name: "매출"
        values: [11, 22, 33, 44]
  - action: set_chart_title
    slide: 1
    chart: 0
    title: "분기 매출"
```
"""


def _req(deck: bytes, instruction: str = "제목·표·차트 고쳐줘") -> EditDocRequest:
    return EditDocRequest(
        content=deck,
        fmt="pptx",
        instruction=instruction,
        lang="ko-KR",
        anthropic_api_key="sk-ant-stub",
    )


class TestPptxSurgicalEdit:
    @pytest.mark.asyncio
    async def test_surgical_edits_apply_in_place(self, monkeypatch):
        import sys

        ed = sys.modules["edit2docs.tools.edit_doc"]
        llm = _PlanLLM(PLAN)
        monkeypatch.setattr(ed, "AnthropicClient", lambda **kw: llm)

        deck = _deck()
        resp = await edit_document(_req(deck))

        assert resp.changed is True
        assert {op["action"] for op in resp.operations} == {
            "set_text", "set_table_cell", "set_chart_data", "set_chart_title"
        }

        prs = Presentation(io.BytesIO(resp.content))
        s1 = prs.slides[0]
        assert "새 제목" in [sh.text_frame.text for sh in s1.shapes if sh.has_text_frame]
        tbl = next(sh.table for sh in s1.shapes if sh.has_table)
        assert tbl.cell(0, 0).text == "바뀐셀"
        # table stayed a real table (still 3 rows, other cells intact)
        assert len(tbl.rows) == 3 and tbl.cell(2, 2).text == "R2C2"
        chart = next(sh.chart for sh in s1.shapes if sh.has_chart)
        assert list(chart.plots[0].categories) == ["Q1", "Q2", "Q3", "Q4"]
        assert list(chart.series[0].values) == [11.0, 22.0, 33.0, 44.0]
        assert chart.chart_title.text_frame.text == "분기 매출"

    @pytest.mark.asyncio
    async def test_outline_shown_to_planner_has_addresses(self, monkeypatch):
        import sys

        ed = sys.modules["edit2docs.tools.edit_doc"]
        llm = _PlanLLM("```reply\n네\n```\n```edit_plan\noperations: []\n```")
        monkeypatch.setattr(ed, "AnthropicClient", lambda **kw: llm)

        await edit_document(_req(_deck(), "이 덱 구성 알려줘"))
        outline = llm.calls[0]
        assert "Deck outline" in outline
        assert "table 3" in outline  # table addressed by shape id
        assert "chart 0" in outline  # chart addressed
        assert "cell (0,0)" in outline  # cells addressable

    @pytest.mark.asyncio
    async def test_question_only_leaves_deck_untouched(self, monkeypatch):
        import sys

        ed = sys.modules["edit2docs.tools.edit_doc"]
        llm = _PlanLLM("```reply\n표지·본문 구성입니다.\n```\n```edit_plan\noperations: []\n```")
        monkeypatch.setattr(ed, "AnthropicClient", lambda **kw: llm)

        deck = _deck()
        resp = await edit_document(_req(deck, "구성이 어때?"))
        assert resp.changed is False
        assert resp.content == deck

    @pytest.mark.asyncio
    async def test_generative_request_signals_needs_svg(self, monkeypatch):
        """A new-slide / full-redesign plan defers to the SVG generator: the
        executor sees needs_svg and hands the turn to tools.edit_deck."""
        import sys

        ed = sys.modules["edit2docs.tools.edit_doc"]
        plan = (
            "```reply\n요약 슬라이드를 새로 추가합니다.\n```\n"
            "```edit_plan\noperations:\n"
            '  - action: add_slide\n    after: 1\n    brief: "요약 슬라이드"\n```'
        )
        monkeypatch.setattr(ed, "AnthropicClient", lambda **kw: _PlanLLM(plan))

        deck = _deck()
        resp = await edit_document(_req(deck, "마지막에 요약 슬라이드 추가해줘"))
        assert resp.needs_svg is True
        assert resp.changed is False
        assert resp.content == deck  # surgical path made no change; SVG path takes over
